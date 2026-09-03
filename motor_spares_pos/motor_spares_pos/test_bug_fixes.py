import tempfile
import unittest
from openpyxl import load_workbook
from pathlib import Path
from unittest.mock import patch
from urllib.error import URLError

from database.db import configure_database
from models.return_refund import Return
from models.sale import Sale
from services.inventory_service import InventoryService
from services.permission_service import PermissionService
from services.sales_service import SalesService
from services.sync_service import SyncService
from services.auth_service import AuthenticationService
from models.product import Product


class BugFixRegressionTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        configure_database(str(Path(self.temp_dir.name) / "pos.db"))
        self.sync_patch = patch(
            "services.sync_service.SyncService.request_background_sync"
        )
        self.sync_patch.start()
        self.product_service = __import__(
            "services.product_service", fromlist=["ProductService"]
        ).ProductService()
        self.inventory = InventoryService()
        self.sales = SalesService()
        product_id = self.product_service.create_product(
            Product(
                barcode="1234567890123",
                part_no="REG-001",
                description="Test part",
                selling_price=10.0,
                quantity_on_hand=10,
            )
        )
        self.product_id = product_id

    def tearDown(self):
        self.sync_patch.stop()
        self.temp_dir.cleanup()

    def test_signed_adjustments(self):
        self.inventory.adjust_stock(self.product_id, 5, "receipt")
        self.assertEqual(
            self.product_service.get_product_by_id(
                self.product_id
            ).quantity_on_hand,
            15,
        )
        movement_id = self.inventory.adjust_stock(
            self.product_id, -5, "correction"
        )
        product = self.product_service.get_product_by_id(self.product_id)
        self.assertEqual(product.quantity_on_hand, 10)
        movement = self.inventory.get_stock_movement(movement_id)
        self.assertEqual(movement.quantity, -5)
        self.assertEqual(movement.previous_quantity, 15)
        self.assertEqual(movement.new_quantity, 10)

    def test_admin_can_permanently_delete_product_without_history(self):
        self.assertTrue(
            self.product_service.delete_product(
                self.product_id, 1, reason="Test cleanup"
            )
        )
        self.assertIsNone(
            self.product_service.get_product_by_id(self.product_id)
        )
        self.assertTrue(
            self.product_service.db.execute_query(
                "SELECT 1 FROM sqlite_master WHERE type='trigger' AND name='prevent_product_physical_delete'"
            )
        )

    def test_archived_product_persists_and_stale_sync_cannot_resurrect_it(
        self,
    ):
        self.assertTrue(
            self.product_service.archive_product(self.product_id, 1)
        )
        self.assertFalse(
            self.product_service.get_product_by_id(self.product_id).active
        )
        sync = SyncService()
        self.assertFalse(
            sync._apply_remote_product(
                {"part_no": "REG-001", "description": "old", "active": True}
            )
        )
        self.assertFalse(
            self.product_service.get_product_by_id(self.product_id).active
        )

    def test_void_restores_stock_once_and_audits_with_valid_schema(self):
        sale = self.sales.create_sale(
            "Customer", user_id=1, cashier_name="Administrator"
        )
        self.sales.add_item_to_sale(sale, self.product_id, 2)
        self.sales.add_payment(sale, "CASH_USD", sale.total, "USD", sale.total)
        self.sales.complete_sale(sale)
        self.assertEqual(
            self.product_service.get_product_by_id(
                self.product_id
            ).quantity_on_hand,
            8,
        )

        self.assertTrue(
            self.sales.void_sale(sale.id, user_id=1, reason="Correction")
        )
        self.assertEqual(
            self.product_service.get_product_by_id(
                self.product_id
            ).quantity_on_hand,
            10,
        )
        movement = self.sales.db.execute_query(
            "SELECT * FROM stock_movements WHERE movement_type='VOID' AND reference=?",
            (sale.invoice_number,),
        )
        self.assertEqual(len(movement), 1)
        audit = self.sales.db.execute_query(
            "SELECT action FROM audit_logs WHERE entity_type='SALE' AND entity_id=?",
            (sale.id,),
        )
        self.assertEqual(audit[0]["action"], "SALE_VOIDED")
        with self.assertRaises(ValueError):
            self.sales.void_sale(sale.id, user_id=1)

    def test_health_requires_success_status(self):
        service = SyncService()
        service.config["server_url"] = "http://sync.example"

        class Response:
            def __init__(self, status):
                self.status = status

            def __enter__(self):
                return self

            def __exit__(self, *_):
                return False

        for status, expected in (
            (200, True),
            (204, True),
            (404, False),
            (500, False),
        ):
            with patch(
                "services.sync_service.urlopen", return_value=Response(status)
            ):
                self.assertEqual(service.is_online(), expected)
        with patch(
            "services.sync_service.urlopen", side_effect=URLError("offline")
        ):
            self.assertFalse(service.is_online())

    def test_role_permission_api_uses_existing_user_permissions_table(self):
        permission = self.sales.db.execute_query(
            "SELECT id FROM permissions WHERE code='VIEW_SALES'"
        )[0]["id"]
        roles = self.sales.db.execute_query(
            "SELECT id FROM roles WHERE name='CASHIER'"
        )
        role_id = roles[0]["id"]
        self.sales.db.execute_update(
            "INSERT INTO users (username, password_hash, full_name, role_id) VALUES (?, ?, ?, ?)",
            ("cashier_fixture", "not-used", "Cashier Fixture", role_id),
        )
        self.assertTrue(
            PermissionService().grant_role_permission(role_id, permission)
        )
        self.assertTrue(
            any(
                p.id == permission
                for p in PermissionService().get_role_permissions(role_id)
            )
        )
        self.assertTrue(
            PermissionService().revoke_role_permission(role_id, permission)
        )

    def test_serialization_does_not_recalculate(self):
        sale = Sale(subtotal=99.0, vat_amount=1.0, total=100.0)
        sale.to_dict()
        self.assertEqual(
            (sale.subtotal, sale.vat_amount, sale.total), (99.0, 1.0, 100.0)
        )
        return_obj = Return(subtotal=50.0, vat_amount=5.0, total_refund=55.0)
        return_obj.to_dict()
        self.assertEqual(
            (
                return_obj.subtotal,
                return_obj.vat_amount,
                return_obj.total_refund,
            ),
            (50.0, 5.0, 55.0),
        )

    def test_admin_can_delete_user_access(self):
        auth = AuthenticationService()
        user_id = auth.create_user(
            "former_employee", "test123", "Former Employee", 3
        )
        self.assertIsNotNone(auth.authenticate("former_employee", "test123"))
        self.assertTrue(auth.delete_user(1, user_id, "Left employment"))
        self.assertIsNone(auth.authenticate("former_employee", "test123"))
        deleted = auth.get_user_by_id(user_id)
        self.assertFalse(deleted.is_active)
        self.assertEqual(auth.get_user_permissions(user_id), [])
        self.assertTrue(
            auth.db.execute_query(
                "SELECT 1 FROM audit_logs WHERE action='USER_DELETED' AND entity_id=?",
                (user_id,),
            )
        )

    def test_customer_matching_by_email_and_name_without_phone(self):
        first = self.sales.find_or_create_customer(
            "Walk In", email="buyer@example.com"
        )
        second = self.sales.find_or_create_customer(
            "Other Name", email="buyer@example.com"
        )
        self.assertEqual(first, second)
        third = self.sales.find_or_create_customer("Walk In")
        fourth = self.sales.find_or_create_customer("Walk In")
        self.assertEqual(third, fourth)

    def test_customer_fields_round_trip_and_sync_payload(self):
        customer_id = self.sales.find_or_create_customer(
            "Harare Buyer", email="buyer@example.com", city="Harare"
        )
        customer = self.sales.get_customer(customer_id)
        self.assertEqual(customer.to_dict()["city"], "Harare")
        self.assertTrue(
            self.sales.update_customer(
                customer_id,
                "Harare Buyer",
                email="new@example.com",
                city="Bulawayo",
            )
        )
        updated = self.sales.get_customer(customer_id)
        self.assertEqual(
            (updated.email, updated.city), ("new@example.com", "Bulawayo")
        )
        queued = self.sales.db.execute_query(
            "SELECT payload FROM sync_queue WHERE entity_type='CUSTOMER' AND entity_id=? AND operation='UPDATE'",
            (customer_id,),
        )[0]
        payload = __import__("json").loads(queued["payload"])
        self.assertEqual(
            (payload["email"], payload["city"]),
            ("new@example.com", "Bulawayo"),
        )

        sale = self.sales.create_sale(
            "Harare Buyer",
            customer_id=customer_id,
            customer_email=updated.email,
            customer_city=updated.city,
            user_id=1,
        )
        self.sales.add_item_to_sale(sale, self.product_id, 1)
        self.sales.add_payment(sale, "CASH_ZIG", sale.total, "ZiG", sale.total)
        self.sales.complete_sale(sale)
        invoice = self.sales.db.execute_query(
            "SELECT customer_email, customer_city FROM invoices WHERE sale_id=?",
            (sale.id,),
        )[0]
        self.assertEqual(
            (invoice["customer_email"], invoice["customer_city"]),
            ("new@example.com", "Bulawayo"),
        )

    def test_existing_customer_without_new_fields_remains_valid(self):
        self.sales.db.execute_update(
            "INSERT INTO customers (name, phone) VALUES (?, ?)",
            ("Legacy", "000"),
        )
        customer = self.sales.get_customer(self.sales.db.get_last_insert_id())
        self.assertEqual((customer.email, customer.city), (None, None))

    def test_inventory_spreadsheet_export_contains_current_inventory(self):
        from services.spreadsheet_service import SpreadsheetService

        path = Path(self.temp_dir.name) / "inventory.xlsx"
        count = SpreadsheetService().export_inventory(str(path))
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(count, 1)
        self.assertEqual(
            rows[1][0:4], ("REG-001", "1234567890123", None, "Test part")
        )

    def test_import_batch_undo_removes_new_product_and_restores_existing(self):
        from openpyxl import Workbook
        from services.spreadsheet_service import SpreadsheetService

        path = Path(self.temp_dir.name) / "wrong_import.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [
                "Part No *",
                "Barcode",
                "Description *",
                "Cost Price",
                "Selling Price",
                "Quantity On Hand",
            ]
        )
        sheet.append(["REG-001", "1234567890123", "Changed part", 4, 20, 25])
        sheet.append(["WRONG-001", "999999999", "Wrong new part", 2, 5, 3])
        workbook.save(path)
        service = SpreadsheetService()
        result = service.import_inventory(str(path), user_id=1)
        self.assertEqual((result.created, result.updated), (1, 1))
        self.assertEqual(
            self.product_service.get_product_by_id(
                self.product_id
            ).selling_price,
            20,
        )
        self.assertTrue(
            service.undo_import(result.import_id, 1)["removed"] == 1
        )
        self.assertEqual(
            self.product_service.get_product_by_id(
                self.product_id
            ).selling_price,
            10,
        )
        self.assertIsNone(
            self.product_service.get_product_by_part_no("WRONG-001")
        )
        with self.assertRaises(ValueError):
            service.undo_import(result.import_id, 1)


if __name__ == "__main__":
    unittest.main()
